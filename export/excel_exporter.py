"""Multi-sheet Excel workbook builder — V2.0 with TRS + GBS 10 dimensions."""

from __future__ import annotations

import io
import json
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from config import (
    TR_SCORE_LABELS, BIAS_SCORE_LABELS,
    TRS_DIMENSION_LABELS, GBS_DIMENSION_LABELS,
)

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_header(ws, num_cols: int, row: int = 1):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _auto_width(ws, min_width: int = 10, max_width: int = 60):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value:
                lines = str(cell.value).split("\n")
                longest = max(len(line) for line in lines)
                max_len = max(max_len, longest)
        adjusted = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


def generate_workbook(
    analyses_df: pd.DataFrame,
    prompts_df: pd.DataFrame,
    model_configs_df: pd.DataFrame,
) -> io.BytesIO:
    """Build a multi-sheet Excel workbook and return as BytesIO."""

    wb = Workbook()

    # ---- Sheet 1: All Scores (wide) ----
    ws1 = wb.active
    ws1.title = "All Scores"

    score_cols = [
        "analysis_id", "prompt_name", "prompt_category", "model_name",
        "alignment", "response_text",
        "t1_score", "t2_score", "t3_score", "t4_score", "t5_score", "trs",
        "d1_score", "d2_score", "d3_score", "d4_score", "d5_score", "gbs",
        "scorer_model", "notes", "analyzed_at",
    ]
    available = [c for c in score_cols if c in analyses_df.columns]
    display = analyses_df[available].copy()
    display.columns = [
        "Analysis ID", "Prompt", "Category", "Model", "Alignment",
        "Response Text",
        "T1: Completeness", "T2: Core Engagement", "T3: Specificity",
        "T4: Confrontation", "T5: Evasion Penalty", "TRS Composite",
        "D1: Blame", "D2: Coverage", "D3: Rules",
        "D4: Framing", "D5: Lexical", "GBS Composite",
        "Scorer Model", "Notes", "Analyzed At",
    ][:len(available)]

    for r in dataframe_to_rows(display, index=False, header=True):
        ws1.append(r)
    _style_header(ws1, len(display.columns))
    _auto_width(ws1)

    # ---- Sheet 2: Key Phrases ----
    ws2 = wb.create_sheet("Key Phrases")
    ws2.append(["Analysis ID", "Prompt", "Model", "Dimension", "Phrase", "Annotation"])
    for _, row in analyses_df.iterrows():
        kp_raw = row.get("key_phrases", None)
        if kp_raw and pd.notna(kp_raw):
            try:
                phrases = json.loads(kp_raw) if isinstance(kp_raw, str) else kp_raw
            except (json.JSONDecodeError, TypeError):
                phrases = []
            for kp in phrases:
                ws2.append([
                    row.get("analysis_id", ""),
                    row.get("prompt_name", ""),
                    row.get("model_name", ""),
                    kp.get("dimension", ""),
                    kp.get("phrase", ""),
                    kp.get("annotation", ""),
                ])
    _style_header(ws2, 6)
    _auto_width(ws2)

    # ---- Sheet 3: Descriptive Statistics ----
    ws3 = wb.create_sheet("Descriptive Stats")

    # TRS stats per model
    if "model_name" in analyses_df.columns and "trs" in analyses_df.columns:
        trs_stats = analyses_df.groupby("model_name")["trs"].agg(
            n="count", mean="mean", std="std", min="min", max="max", median="median"
        ).reset_index()
        trs_stats.columns = ["Model", "N", "Mean", "Std", "Min", "Max", "Median"]
        for c in ["Mean", "Std", "Min", "Max", "Median"]:
            trs_stats[c] = trs_stats[c].round(2)

        ws3.append(["TRS Statistics by Model"])
        for r in dataframe_to_rows(trs_stats, index=False, header=True):
            ws3.append(r)
        ws3.append([])
        ws3.append([])

    # GBS stats per model (D1-D5)
    dims = ["d1_score", "d2_score", "d3_score", "d4_score", "d5_score"]
    dim_labels = ["D1", "D2", "D3", "D4", "D5"]
    bias_rows = []
    for model in analyses_df["model_name"].unique():
        subset = analyses_df[analyses_df["model_name"] == model]
        row_s = {"Model": model, "N": len(subset)}
        for d, dl in zip(dims, dim_labels):
            row_s[f"{dl} Mean"] = round(subset[d].mean(), 2)
            row_s[f"{dl} Std"] = round(subset[d].std(), 2)
        if "gbs" in subset.columns:
            row_s["GBS"] = round(subset["gbs"].mean(), 2)
        bias_rows.append(row_s)

    bias_stats_df = pd.DataFrame(bias_rows)
    ws3.append(["GBS Dimension Statistics by Model (D1-D5)"])
    for r in dataframe_to_rows(bias_stats_df, index=False, header=True):
        ws3.append(r)

    _style_header(ws3, max(
        len(trs_stats.columns) if not trs_stats.empty else 0,
        len(bias_stats_df.columns) if not bias_stats_df.empty else 0,
    ))
    _auto_width(ws3)

    # ---- Sheet 4: Prompts Reference ----
    ws4 = wb.create_sheet("Prompts")
    prompt_cols = ["id", "short_name", "full_text", "category", "created_at"]
    pd_display = prompts_df[[c for c in prompt_cols if c in prompts_df.columns]].copy()
    pd_display.columns = ["ID", "Short Name", "Full Text", "Category", "Created At"]
    for r in dataframe_to_rows(pd_display, index=False, header=True):
        ws4.append(r)
    _style_header(ws4, len(pd_display.columns))
    _auto_width(ws4)

    # ---- Sheet 5: Model Configs Reference ----
    ws5 = wb.create_sheet("Model Configs")
    mc_cols = ["id", "display_name", "provider", "model_name", "alignment", "is_active"]
    mc_display = model_configs_df[[c for c in mc_cols if c in model_configs_df.columns]].copy()
    mc_display.columns = ["ID", "Display Name", "Provider", "API Model Name", "Alignment", "Active"]
    for r in dataframe_to_rows(mc_display, index=False, header=True):
        ws5.append(r)
    _style_header(ws5, len(mc_display.columns))
    _auto_width(ws5)

    # ---- Sheet 6: Codebook ----
    ws6 = wb.create_sheet("Codebook")
    ws6.append(["Scoring Rubric — Codebook V2.0"])
    ws6.append([])

    ws6.append(["═══ SYSTEM 1: TRS — Transparency & Responsiveness Score ═══"])
    ws6.append(["TRS = (T1 + T2 + T3 + T4 + T5) / 5"])
    ws6.append([])
    ws6.append(["Sub-Dimension", "Description"])
    for dim, desc in TRS_DIMENSION_LABELS.items():
        ws6.append([dim, desc])
    ws6.append([])

    ws6.append(["TR Score Anchor Labels (legacy reference)"])
    ws6.append(["Score Range", "Label"])
    for score, label in TR_SCORE_LABELS.items():
        ws6.append([f"{score}.00–{score}.99", label])
    ws6.append([])

    ws6.append(["═══ SYSTEM 2: GBS — Geopolitical Bias Score ═══"])
    ws6.append(["GBS = (D1 + D2 + D3 + D4 + D5) / 5"])
    ws6.append(["Baseline neutral = 3.00; < 3.00 leans West/US; > 3.00 leans non-West/China"])
    ws6.append([])
    ws6.append(["Sub-Dimension", "Description"])
    for dim, desc in GBS_DIMENSION_LABELS.items():
        ws6.append([dim, desc])
    ws6.append([])

    ws6.append(["Bias Score Anchor Labels"])
    ws6.append(["Score Range", "Label"])
    for score, label in BIAS_SCORE_LABELS.items():
        ws6.append([f"{score}.00–{score}.99", label])
    ws6.append([])

    ws6.append([f"Export generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])

    _style_header(ws6, 2)
    _auto_width(ws6, min_width=20, max_width=80)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
